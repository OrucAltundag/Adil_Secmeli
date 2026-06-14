# -*- coding: utf-8 -*-
"""Otomatik üretim hattı (auto pipeline).

Kullanıcı talebi: "Sistemin otomatik çalışır durumda olması veya manuel
çalıştırılmasını aktif edilebilmesi... kriter girdi işlemleri tamamlandığı gibi
sistem yeni müfredatları, ders önerilerini hazırlasın, paylaş. Hatta bir sürü
şeyi de Excel olarak oluşturup paylaş."

Bu modül:
- config.json içindeki `auto_pipeline_enabled` bayrağını okur/yazar (otomatik/manuel).
- Kriterleri (her iki dönem) tamamlanmış fakülteler için sonraki yıl müfredatını
  üretir (generate_next_year_curricula üzerinden).
- Üretilen yıl için ders önerisi (müfredat + havuz) raporunu Excel'e döker.
- Çalıştırmaları auto_pipeline_runs tablosuna loglar.

Tasarım: yan etkisi olan tek giriş noktası `run_auto_pipeline`. Hem manuel butondan
hem de kriter importu sonrası tetikleyiciden çağrılabilir.
"""

from __future__ import annotations

import json
import os
import sqlite3
from datetime import datetime
from typing import Any

from app.core.config import resolve_sqlite_db_path

DEFAULT_EXPORT_DIR = "exports"
AUTO_PIPELINE_CONFIG_KEY = "auto_pipeline_enabled"


class _SqliteReportAdapter:
    """build_report_snapshot'ın beklediği minimal db arayüzü (run_sql + conn).

    Global SQLAlchemy engine'e bağlanmadan, verilen db_path üzerinde doğrudan
    çalışır. Böylece auto-pipeline herhangi bir veritabanı yolu için güvenle
    rapor üretir.
    """

    def __init__(self, db_path: str) -> None:
        self.conn = sqlite3.connect(db_path)
        self.conn.row_factory = sqlite3.Row

    def run_sql(self, query: str, params=None):
        cur = self.conn.cursor()
        cur.execute(query, tuple(params) if params else ())
        if query.strip().lower().startswith("select"):
            cols = [d[0] for d in cur.description] if cur.description else []
            rows = [tuple(r) for r in cur.fetchall()]
            return cols, rows
        self.conn.commit()
        return [], []

    def close(self) -> None:
        try:
            self.conn.close()
        except Exception:
            pass


# --------------------------------------------------------------------- config
def is_auto_pipeline_enabled(config_path: str = "config.json") -> bool:
    """Otomatik mod açık mı? Önce ortam değişkeni, sonra config.json."""
    raw_env = os.getenv("ENABLE_AUTO_PIPELINE")
    if raw_env is not None:
        return str(raw_env).strip().lower() in {"1", "true", "yes", "on"}
    try:
        with open(config_path, encoding="utf-8") as fh:
            cfg = json.load(fh)
        return bool(cfg.get(AUTO_PIPELINE_CONFIG_KEY, False))
    except Exception:
        return False


def set_auto_pipeline_enabled(enabled: bool, config_path: str = "config.json") -> bool:
    """Otomatik modu açar/kapatır ve config.json'a yazar. Yeni değeri döndürür."""
    cfg: dict[str, Any] = {}
    try:
        if os.path.exists(config_path):
            with open(config_path, encoding="utf-8") as fh:
                cfg = json.load(fh)
    except Exception:
        cfg = {}
    cfg[AUTO_PIPELINE_CONFIG_KEY] = bool(enabled)
    with open(config_path, "w", encoding="utf-8") as fh:
        json.dump(cfg, fh, ensure_ascii=False, indent=2)
    return bool(enabled)


# ---------------------------------------------------------------- schema/log
def ensure_auto_pipeline_schema(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS auto_pipeline_runs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            started_at TEXT,
            finished_at TEXT,
            trigger TEXT,
            source_year INTEGER,
            generated_year INTEGER,
            faculty_scope TEXT,
            processed_count INTEGER DEFAULT 0,
            skipped_count INTEGER DEFAULT 0,
            error_count INTEGER DEFAULT 0,
            export_path TEXT,
            summary_json TEXT
        )
        """
    )
    conn.commit()


def _log_run(conn: sqlite3.Connection, payload: dict[str, Any]) -> int:
    ensure_auto_pipeline_schema(conn)
    cur = conn.cursor()
    cur.execute(
        """
        INSERT INTO auto_pipeline_runs
            (started_at, finished_at, trigger, source_year, generated_year,
             faculty_scope, processed_count, skipped_count, error_count,
             export_path, summary_json)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            payload.get("started_at"),
            payload.get("finished_at"),
            payload.get("trigger"),
            payload.get("source_year"),
            payload.get("generated_year"),
            payload.get("faculty_scope"),
            int(payload.get("processed_count") or 0),
            int(payload.get("skipped_count") or 0),
            int(payload.get("error_count") or 0),
            payload.get("export_path"),
            json.dumps(payload.get("summary") or {}, ensure_ascii=False, default=str),
        ),
    )
    conn.commit()
    return int(cur.lastrowid or 0)


# --------------------------------------------------------------- ana akış
def _faculties(conn: sqlite3.Connection, faculty_id: int | None) -> list[tuple[int, str]]:
    cur = conn.cursor()
    if faculty_id is not None:
        cur.execute("SELECT fakulte_id, ad FROM fakulte WHERE fakulte_id = ?", (int(faculty_id),))
    else:
        cur.execute("SELECT fakulte_id, ad FROM fakulte ORDER BY fakulte_id")
    return [(int(r[0]), str(r[1] or "")) for r in cur.fetchall()]


def _active_terms(conn: sqlite3.Connection, faculty_id: int, year: int) -> list[str]:
    """Fakültenin o yıl müfredatında bulunan dönemler ('G'/'B')."""
    cur = conn.cursor()
    cur.execute(
        """
        SELECT DISTINCT LOWER(SUBSTR(TRIM(COALESCE(m.donem, '')), 1, 1))
        FROM mufredat m
        JOIN bolum b ON b.bolum_id = m.bolum_id
        WHERE b.fakulte_id = ? AND m.akademik_yil = ?
          AND COALESCE(TRIM(m.donem), '') <> ''
        """,
        (int(faculty_id), int(year)),
    )
    keys = {str(r[0]) for r in cur.fetchall() if r and r[0] in ("g", "b")}
    # Hiç bulunmazsa varsayılan Güz.
    return sorted(keys) or ["g"]


def run_auto_pipeline(
    db_path: str,
    source_year: int,
    faculty_id: int | None = None,
    export_dir: str = DEFAULT_EXPORT_DIR,
    trigger: str = "manual",
    export: bool = True,
) -> dict[str, Any]:
    """Kriterleri tamam fakülteler için sonraki yıl müfredatını üretir ve Excel'e döker.

    Returns: {ok, generated_year, processed[], skipped[], errors[], export_path, ...}
    """
    from app.services.calculation import generate_next_year_curricula

    started_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    resolved = resolve_sqlite_db_path(db_path)
    if not resolved.exists():
        return {"ok": False, "error": f"Veritabani bulunamadi: {resolved}"}

    source_year = int(source_year)
    generated_year = source_year + 1

    conn = sqlite3.connect(str(resolved))
    conn.row_factory = sqlite3.Row
    processed: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    try:
        faculties = _faculties(conn, faculty_id)
        for fid, fname in faculties:
            terms = _active_terms(conn, fid, source_year)
            faculty_ok_any = False
            for term_key in terms:
                donem = "B" if term_key == "b" else "G"
                try:
                    result = generate_next_year_curricula(
                        db_path=str(resolved),
                        fakulte_id=fid,
                        akademik_yil=source_year,
                        donem=donem,
                    )
                except Exception as exc:  # noqa: BLE001
                    errors.append({"fakulte": fname, "donem": donem, "error": str(exc)})
                    continue
                if result.get("ok"):
                    faculty_ok_any = True
                    processed.append({"fakulte": fname, "faculty_id": fid, "donem": donem})
                else:
                    skipped.append(
                        {
                            "fakulte": fname,
                            "faculty_id": fid,
                            "donem": donem,
                            "reason": result.get("error", "Bilinmeyen"),
                            "blocked_terms": result.get("blocked_terms"),
                        }
                    )
            if not faculty_ok_any and not any(s["faculty_id"] == fid for s in skipped):
                skipped.append({"fakulte": fname, "faculty_id": fid, "reason": "İşlenecek dönem yok"})
    finally:
        conn.close()

    export_path = None
    if export and processed:
        try:
            export_path = export_recommendations_excel(
                db_path=str(resolved),
                year=generated_year,
                export_dir=export_dir,
                faculty_ids=sorted({int(p["faculty_id"]) for p in processed}),
            )
        except Exception as exc:  # noqa: BLE001
            errors.append({"stage": "excel_export", "error": str(exc)})

    finished_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    summary = {
        "ok": bool(processed) and not errors,
        "trigger": trigger,
        "source_year": source_year,
        "generated_year": generated_year,
        "faculty_scope": "all" if faculty_id is None else str(faculty_id),
        "processed": processed,
        "skipped": skipped,
        "errors": errors,
        "export_path": export_path,
        "started_at": started_at,
        "finished_at": finished_at,
    }

    # Çalıştırma logu (best-effort)
    try:
        log_conn = sqlite3.connect(str(resolved))
        _log_run(
            log_conn,
            {
                **summary,
                "processed_count": len(processed),
                "skipped_count": len(skipped),
                "error_count": len(errors),
                "summary": summary,
            },
        )
        log_conn.close()
    except Exception:
        pass

    return summary


# ---------------------------------------------------------------- Excel çıktı
def export_recommendations_excel(
    db_path: str,
    year: int,
    export_dir: str = DEFAULT_EXPORT_DIR,
    faculty_ids: list[int] | None = None,
) -> str:
    """Üretilen yıl için ders önerisi (müfredat + havuz) raporunu Excel'e döker.

    Her fakülte için müfredat ve havuz satırlarını ayrı sayfalarda toplar; ayrıca
    bir özet sayfası ekler. Döküm yolu döner.
    """
    import pandas as pd
    from openpyxl.styles import Font

    from app.services.reporting_service import build_report_snapshot

    resolved = resolve_sqlite_db_path(db_path)
    db = _SqliteReportAdapter(str(resolved))
    try:
        conn = db.conn
        cur = conn.cursor()
        if faculty_ids:
            placeholders = ",".join("?" for _ in faculty_ids)
            cur.execute(
                f"SELECT fakulte_id, ad FROM fakulte WHERE fakulte_id IN ({placeholders}) ORDER BY ad",
                tuple(int(f) for f in faculty_ids),
            )
        else:
            cur.execute("SELECT fakulte_id, ad FROM fakulte ORDER BY ad")
        faculties = [(int(r[0]), str(r[1] or "")) for r in cur.fetchall()]

        os.makedirs(export_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_path = os.path.join(export_dir, f"ders_onerisi_{year}_{timestamp}.xlsx")

        summary_rows: list[dict[str, Any]] = []
        mufredat_all: list[dict[str, Any]] = []
        havuz_all: list[dict[str, Any]] = []

        for fid, fname in faculties:
            for term_key, term_label in (("G", "Güz"), ("B", "Bahar")):
                try:
                    snap = build_report_snapshot(
                        db=db, faculty_id=fid, faculty_name=fname, year=int(year), term=term_label,
                    )
                except Exception:
                    continue
                pool_rows = snap.get("pool_rows") or []
                if not pool_rows:
                    continue
                stats = snap.get("stats") or {}
                summary_rows.append(
                    {
                        "Fakülte": fname,
                        "Dönem": term_label,
                        "Yıl": int(year),
                        "Müfredattaki Ders": stats.get("chosen_count", 0),
                        "Dinlenmede": stats.get("rest_count", 0),
                        "İptal": stats.get("cancelled_count", 0),
                        "Havuz Toplam": stats.get("total", 0),
                        "Ortalama Skor": stats.get("avg_score"),
                    }
                )
                for row in pool_rows:
                    record = {
                        "Fakülte": fname,
                        "Dönem": term_label,
                        "Ders ID": row.get("ders_id"),
                        "Ders Adı": row.get("ders_adi"),
                        "Kesinleşme Puanı": row.get("skor"),
                        "Sayaç": row.get("sayac"),
                        "Durum": row.get("statu"),
                        "Skor Kaynağı": row.get("kaynak"),
                    }
                    havuz_all.append(record)
                    # Müfredata seçilenler (Mufredatta) ayrı sayfada da listelensin.
                    if str(row.get("statu", "")).lower().startswith("müfredat") or str(row.get("statu", "")).lower().startswith("mufredat"):
                        mufredat_all.append(record)

        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            # Özet
            df_summary = pd.DataFrame(summary_rows) if summary_rows else pd.DataFrame(
                [{"Bilgi": "Üretilen veri bulunamadı"}]
            )
            df_summary.to_excel(writer, sheet_name="Özet", index=False)
            # Müfredat önerisi
            df_muf = pd.DataFrame(mufredat_all) if mufredat_all else pd.DataFrame(
                [{"Bilgi": "Müfredat önerisi yok"}]
            )
            df_muf.to_excel(writer, sheet_name="Müfredat Önerisi", index=False)
            # Havuz durumu
            df_hav = pd.DataFrame(havuz_all) if havuz_all else pd.DataFrame(
                [{"Bilgi": "Havuz verisi yok"}]
            )
            df_hav.to_excel(writer, sheet_name="Havuz Durumu", index=False)

            # Başlık satırlarını kalın yap
            for sheet_name in ("Özet", "Müfredat Önerisi", "Havuz Durumu"):
                ws = writer.sheets[sheet_name]
                for cell in ws[1]:
                    cell.font = Font(bold=True)

        return out_path
    finally:
        db.close()


def list_recent_auto_runs(db_path: str, limit: int = 20) -> list[dict[str, Any]]:
    resolved = resolve_sqlite_db_path(db_path)
    if not resolved.exists():
        return []
    conn = sqlite3.connect(str(resolved))
    conn.row_factory = sqlite3.Row
    try:
        ensure_auto_pipeline_schema(conn)
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, started_at, finished_at, trigger, source_year, generated_year,
                   faculty_scope, processed_count, skipped_count, error_count, export_path
            FROM auto_pipeline_runs
            ORDER BY id DESC
            LIMIT ?
            """,
            (int(limit),),
        )
        return [dict(row) for row in cur.fetchall()]
    finally:
        conn.close()
