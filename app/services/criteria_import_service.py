from __future__ import annotations

import json
import os
import sqlite3
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from typing import Any

import pandas as pd
from openpyxl.styles import Font

from app.core.config import resolve_sqlite_db_path
from app.db.schema_compat import ensure_criteria_import_schema, ensure_reporting_schema
from app.db.sqlite_connection import connect_sqlite, is_database_locked_error
from app.services.course_matcher import (
    CourseCandidate,
    load_faculty_course_candidates,
    match_course_row,
    normalize_course_key,
    normalize_course_text,
)
from app.services.course_type import build_elective_predicate
from app.services.import_audit_service import (
    calculate_row_hash,
    create_import_batch,
    extract_excel_metadata,
    link_source_import,
    mark_batch_failed_by_path,
    record_import_issue,
    update_import_status,
)
from app.services.import_diff_service import recalculate_import_diff
from app.services.import_impact_service import recalculate_import_impact
from app.services.import_lineage_service import record_value_source
from app.services.import_quality_service import evaluate_import_quality
from app.services.yearly_workflow import mark_criteria_status

CRITERIA_TEMPLATE_VERSION = "criteria-import-v1"
CRITERIA_TEMPLATE_SHEET_NAME = "Kriter Veri Giris Sablonu"
FACULTY_SCOPE_LABEL = "Fakulte Geneli"


@dataclass
class CriteriaRow:
    row_no: int
    ders_kodu: str | None
    ders_adi: str | None
    toplam_ogrenci: int | None
    gecen_ogrenci: int | None
    basari_ortalamasi: float | None
    kontenjan: int | None
    kayitli_ogrenci: int | None
    aciklama: str | None = None
    fakulte_adi: str | None = None
    bolum_adi: str | None = None
    yil: int | None = None
    donem: str | None = None


@dataclass
class CriteriaImportRowResult:
    row_no: int
    ders_kodu: str | None
    ders_adi: str | None
    toplam_ogrenci: int
    gecen_ogrenci: int
    basari_ortalamasi: float
    kontenjan: int
    kayitli_ogrenci: int
    aciklama: str | None = None
    matched_ders_id: int | None = None
    match_method: str | None = None
    row_status: str = "matched"
    error_message: str | None = None
    raw_fakulte: str | None = None
    raw_bolum: str | None = None
    raw_yil: int | None = None
    raw_donem: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


def _now_utc() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def normalize_term_label(value: str | None) -> str:
    raw = str(value or "").strip().lower()
    if raw.startswith("b"):
        return "Bahar"
    return "Güz"


def term_key(value: str | None) -> str:
    return "b" if normalize_term_label(value) == "Bahar" else "g"


def normalize_department_scope_name(value: str | None) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    if normalize_course_text(text) == normalize_course_text(FACULTY_SCOPE_LABEL):
        return None
    return text


def _find_col(columns: list[str], *candidates: str) -> str | None:
    normalized = {normalize_course_text(col): col for col in columns}
    for cand in candidates:
        key = normalize_course_text(cand)
        if key in normalized:
            return normalized[key]
    return None


def _parse_year(value: Any) -> int | None:
    try:
        if value is None or pd.isna(value):
            return None
        year = int(float(value))
    except (TypeError, ValueError):
        text = str(value or "").strip()
        digits = "".join(ch for ch in text if ch.isdigit())
        if len(digits) >= 4:
            try:
                year = int(digits[:4])
            except ValueError:
                return None
        else:
            return None
    return year if 1900 <= year <= 2100 else None


def _safe_int(value: Any) -> int | None:
    try:
        if value is None or pd.isna(value):
            return None
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _safe_float(value: Any) -> float | None:
    try:
        if value is None or pd.isna(value):
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def _clean_text(value: Any) -> str | None:
    if value is None or pd.isna(value):
        return None
    text = str(value).strip()
    return text or None


def _excel_column_letter(index_1_based: int) -> str:
    if index_1_based < 1:
        raise ValueError("Excel kolon indeksi 1 veya daha buyuk olmali.")
    out = ""
    current = int(index_1_based)
    while current > 0:
        current, remainder = divmod(current - 1, 26)
        out = chr(65 + remainder) + out
    return out


def _is_summary_row(ders_kodu: str | None, ders_adi: str | None) -> bool:
    if _clean_text(ders_kodu):
        return False
    normalized_name = normalize_course_key(ders_adi)
    return normalized_name in {"toplam", "geneltoplam"}


def _read_meta_sheet(xls: pd.ExcelFile) -> dict[str, Any]:
    meta_sheet = next((name for name in xls.sheet_names if normalize_course_text(str(name)) == "meta"), None)
    if not meta_sheet:
        return {}
    df = pd.read_excel(xls, sheet_name=meta_sheet)
    if df.empty:
        return {}
    df.columns = [str(col).strip() for col in df.columns]
    row = df.iloc[0].to_dict()
    return {
        "fakulte_adi": _clean_text(row.get(_find_col(list(df.columns), "fakulte_adi", "fakulte", "faculty"))),
        "bolum_adi": _clean_text(row.get(_find_col(list(df.columns), "bolum_adi", "bolum", "department"))),
        "yil": _parse_year(row.get(_find_col(list(df.columns), "yil", "akademik_yil", "year"))),
        "donem": normalize_term_label(row.get(_find_col(list(df.columns), "donem", "term", "semester"))),
        "aciklama": _clean_text(row.get(_find_col(list(df.columns), "aciklama", "not", "notes"))),
    }


def parse_criteria_excel(excel_path: str) -> dict[str, Any]:
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Kriter dosyasi bulunamadi: {excel_path}")

    xls = pd.ExcelFile(excel_path)
    meta = _read_meta_sheet(xls)
    data_sheet = next(
        (
            name
            for name in xls.sheet_names
            if normalize_course_text(str(name)) in {"kriter", "kriterler", "criteria", "criterion"}
        ),
        None,
    )
    if data_sheet is None:
        non_meta = [name for name in xls.sheet_names if normalize_course_text(str(name)) != "meta"]
        if not non_meta:
            raise ValueError("Kriter veri sayfasi bulunamadi.")
        data_sheet = non_meta[0]

    df = pd.read_excel(xls, sheet_name=data_sheet)
    df.columns = [str(col).strip() for col in df.columns]
    columns = list(df.columns)

    col_code = _find_col(columns, "ders_kodu", "ders kodu", "kod", "course_code")
    col_name = _find_col(columns, "ders_adi", "ders adi", "course_name")
    col_total = _find_col(columns, "toplam_ogrenci", "toplam ogrenci", "dersi_alan_toplam_ogrenci")
    col_pass = _find_col(columns, "gecen_ogrenci", "gecen ogrenci", "dersi_gecen_ogrenci")
    col_avg = _find_col(columns, "basari_ortalamasi", "ortalama_not", "ortalama")
    col_quota = _find_col(columns, "kontenjan", "ders_kontenjani")
    col_enrolled = _find_col(columns, "kayitli_ogrenci", "kayitli ogrenci", "talep_sayisi")
    col_note = _find_col(columns, "aciklama", "not")
    col_faculty = _find_col(columns, "fakulte_adi", "fakulte", "faculty")
    col_department = _find_col(columns, "bolum_adi", "bolum", "department")
    col_year = _find_col(columns, "yil", "akademik_yil", "year")
    col_term = _find_col(columns, "donem", "term", "semester")

    required_numeric_cols = [col_total, col_pass, col_avg, col_quota]
    if not all(required_numeric_cols):
        raise ValueError("Gerekli kolonlar bulunamadi: toplam_ogrenci, gecen_ogrenci, basari_ortalamasi, kontenjan")
    if not (col_code or col_name):
        raise ValueError("Ders tanimlayici gerekli: ders_kodu veya ders_adi")

    rows: list[CriteriaRow] = []
    warnings: list[str] = []
    for idx, row in df.iterrows():
        ders_kodu = _clean_text(row.get(col_code)) if col_code else None
        ders_adi = _clean_text(row.get(col_name)) if col_name else None
        if not ders_kodu and not ders_adi:
            continue
        if _is_summary_row(ders_kodu=ders_kodu, ders_adi=ders_adi):
            continue

        toplam = _safe_int(row.get(col_total))
        gecen = _safe_int(row.get(col_pass))
        ortalama = _safe_float(row.get(col_avg))
        kontenjan = _safe_int(row.get(col_quota))
        kayitli = _safe_int(row.get(col_enrolled)) if col_enrolled else None
        if kayitli is None:
            kayitli = toplam

        criteria_row = CriteriaRow(
            row_no=int(str(idx)) + 2,
            ders_kodu=ders_kodu,
            ders_adi=ders_adi,
            toplam_ogrenci=toplam,
            gecen_ogrenci=gecen,
            basari_ortalamasi=ortalama,
            kontenjan=kontenjan,
            kayitli_ogrenci=kayitli,
            aciklama=_clean_text(row.get(col_note)) if col_note else None,
            fakulte_adi=_clean_text(row.get(col_faculty)) if col_faculty else None,
            bolum_adi=_clean_text(row.get(col_department)) if col_department else None,
            yil=_parse_year(row.get(col_year)) if col_year else None,
            donem=normalize_term_label(row.get(col_term)) if col_term else None,
        )
        if any(
            value is None
            for value in (
                criteria_row.toplam_ogrenci,
                criteria_row.gecen_ogrenci,
                criteria_row.basari_ortalamasi,
                criteria_row.kontenjan,
            )
        ):
            warnings.append(f"Satir {criteria_row.row_no}: zorunlu sayisal alanlardan biri bos veya gecersiz.")
        rows.append(criteria_row)

    return {
        "meta": meta,
        "rows": rows,
        "warnings": warnings,
        "sheet_name": data_sheet,
        "template_version": CRITERIA_TEMPLATE_VERSION,
    }


def validate_criteria_rows(
    rows: list[CriteriaRow],
    faculty_name: str | None = None,
    department_name: str | None = None,
    year: int | None = None,
    term: str | None = None,
) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []

    if not rows:
        errors.append("Belge icinde aktarilabilir kriter satiri yok.")
        return {"ok": False, "errors": errors, "warnings": warnings}

    seen_document_keys: dict[str, int] = {}
    normalized_department = normalize_department_scope_name(department_name)
    normalized_term = normalize_term_label(term)

    for row in rows:
        if not row.ders_kodu and not row.ders_adi:
            errors.append(f"Satir {row.row_no}: ders_kodu veya ders_adi zorunlu.")
            continue

        if row.toplam_ogrenci is None or row.gecen_ogrenci is None or row.basari_ortalamasi is None or row.kontenjan is None:
            errors.append(f"Satir {row.row_no}: zorunlu kriter alanlari eksik.")
            continue

        if row.kayitli_ogrenci is None:
            row.kayitli_ogrenci = row.toplam_ogrenci

        if row.toplam_ogrenci < 0 or row.gecen_ogrenci < 0 or row.kontenjan < 0 or int(row.kayitli_ogrenci) < 0:
            errors.append(f"Satir {row.row_no}: sayisal alanlar negatif olamaz.")
        if row.gecen_ogrenci > row.toplam_ogrenci:
            errors.append(f"Satir {row.row_no}: gecen_ogrenci toplam_ogrenci degerini asamaz.")
        if faculty_name and row.fakulte_adi and normalize_course_text(row.fakulte_adi) != normalize_course_text(faculty_name):
            errors.append(
                f"Satir {row.row_no}: belge fakultesi '{row.fakulte_adi}' secili fakulte '{faculty_name}' ile uyusmuyor."
            )
        if normalized_department and row.bolum_adi and normalize_course_text(row.bolum_adi) != normalize_course_text(normalized_department):
            errors.append(
                f"Satir {row.row_no}: belge bolumu '{row.bolum_adi}' secili bolum '{normalized_department}' ile uyusmuyor."
            )
        if year is not None and row.yil is not None and int(row.yil) != int(year):
            errors.append(f"Satir {row.row_no}: belge yili '{row.yil}' secili yil '{year}' ile uyusmuyor.")
        if term is not None and row.donem is not None and term_key(row.donem) != term_key(normalized_term):
            errors.append(
                f"Satir {row.row_no}: belge donemi '{row.donem}' secili donem '{normalized_term}' ile uyusmuyor."
            )

        dedupe_key = normalize_course_text(row.ders_kodu) or f"ad:{normalize_course_key(row.ders_adi)}"
        if dedupe_key:
            if dedupe_key in seen_document_keys:
                errors.append(
                    f"Belgede ayni ders birden fazla kez geciyor: satir {seen_document_keys[dedupe_key]} ve {row.row_no}."
                )
            else:
                seen_document_keys[dedupe_key] = row.row_no

    return {
        "ok": len(errors) == 0,
        "errors": errors,
        "warnings": warnings,
    }


def _resolve_faculty_name(cur: sqlite3.Cursor, faculty_id: int) -> str | None:
    cur.execute("SELECT ad FROM fakulte WHERE fakulte_id = ? LIMIT 1", (int(faculty_id),))
    row = cur.fetchone()
    return str(row[0] or "") if row else None


def _resolve_department_name(cur: sqlite3.Cursor, department_id: int | None) -> str | None:
    if department_id is None:
        return None
    cur.execute("SELECT ad FROM bolum WHERE bolum_id = ? LIMIT 1", (int(department_id),))
    row = cur.fetchone()
    return str(row[0] or "") if row else None


def _table_exists(cur: sqlite3.Cursor, table_name: str) -> bool:
    cur.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=? LIMIT 1",
        (str(table_name),),
    )
    return cur.fetchone() is not None


def _elective_predicate(cur: sqlite3.Cursor, alias: str) -> str:
    try:
        predicate = build_elective_predicate(cur=cur, alias=alias)
    except Exception:
        predicate = "1=1"
    return "1=1" if predicate == "0=1" else predicate


def _get_scope_courses(
    cur: sqlite3.Cursor,
    faculty_id: int,
    year: int,
    term: str,
    department_id: int | None = None,
) -> list[dict[str, Any]]:
    predicate = _elective_predicate(cur=cur, alias="d")
    params: list[Any] = [int(faculty_id), int(year), term_key(term)]
    department_clause = ""
    if department_id is not None:
        department_clause = "AND b.bolum_id = ?"
        params.append(int(department_id))

    cur.execute(
        f"""
        SELECT DISTINCT
            d.ders_id,
            NULLIF(TRIM(COALESCE(d.kod, '')), '') AS ders_kodu,
            COALESCE(NULLIF(TRIM(d.ad), ''), 'Ders ' || d.ders_id) AS ders_adi,
            b.bolum_id,
            b.ad AS bolum_adi
        FROM mufredat m
        JOIN bolum b ON b.bolum_id = m.bolum_id
        JOIN mufredat_ders md ON md.mufredat_id = m.mufredat_id
        JOIN ders d ON d.ders_id = md.ders_id
        WHERE b.fakulte_id = ?
          AND m.akademik_yil = ?
          AND LOWER(SUBSTR(TRIM(COALESCE(m.donem, '')), 1, 1)) = ?
          {department_clause}
          AND {predicate}
        ORDER BY b.ad, d.ad, d.ders_id
        """,
        tuple(params),
    )
    return [
        {
            "ders_id": int(row[0]),
            "ders_kodu": _clean_text(row[1]),
            "ders_adi": str(row[2] or "").strip(),
            "bolum_id": int(row[3]) if row[3] is not None else None,
            "bolum_adi": str(row[4] or "").strip() or None,
        }
        for row in cur.fetchall()
        if row and row[0] is not None and str(row[2] or "").strip()
    ]


def load_criteria_template_context(
    db_path: str,
    faculty_id: int,
    year: int,
    term: str,
    department_id: int | None = None,
) -> dict[str, Any]:
    resolved_db_path = resolve_sqlite_db_path(db_path)
    if not resolved_db_path.exists():
        raise FileNotFoundError("Veritabani bulunamadi.")

    conn = connect_sqlite(str(resolved_db_path))
    try:
        try:
            ensure_reporting_schema(conn)
        except sqlite3.OperationalError as exc:
            if not is_database_locked_error(exc):
                raise
            conn.rollback()
        cur = conn.cursor()
        faculty_name = _resolve_faculty_name(cur, int(faculty_id))
        if not faculty_name:
            raise ValueError("Secili fakulte bulunamadi.")
        department_name = _resolve_department_name(cur, int(department_id)) if department_id is not None else None
        courses = _get_scope_courses(
            cur=cur,
            faculty_id=int(faculty_id),
            year=int(year),
            term=normalize_term_label(term),
            department_id=int(department_id) if department_id is not None else None,
        )
        if not courses:
            raise ValueError("Secili kapsam icin mufredatta uygun ders bulunamadi.")
        return {
            "faculty_name": faculty_name,
            "department_name": department_name,
            "year": int(year),
            "term": normalize_term_label(term),
            "courses": courses,
        }
    finally:
        conn.close()


def write_criteria_template_excel(
    target_path: str,
    faculty_name: str | None = None,
    department_name: str | None = None,
    year: int | None = None,
    term: str | None = None,
    db_path: str | None = None,
    faculty_id: int | None = None,
    department_id: int | None = None,
) -> str:
    template_courses: list[dict[str, Any]] = []
    if db_path and faculty_id is not None and year is not None and term is not None:
        context = load_criteria_template_context(
            db_path=db_path,
            faculty_id=int(faculty_id),
            year=int(year),
            term=normalize_term_label(term),
            department_id=int(department_id) if department_id is not None else None,
        )
        faculty_name = str(context.get("faculty_name") or faculty_name or "")
        department_name = str(context.get("department_name") or department_name or "") or None
        template_courses = list(context.get("courses") or [])

    faculty_text = faculty_name or "Ornek Fakultesi"
    department_text = normalize_department_scope_name(department_name)
    year_value = int(year or datetime.now().year)
    term_value = normalize_term_label(term)

    if template_courses:
        rows = [
            {
                "fakulte_adi": faculty_text,
                "bolum_adi": course.get("bolum_adi") or department_text or FACULTY_SCOPE_LABEL,
                "yil": year_value,
                "donem": term_value,
                "ders_kodu": course.get("ders_kodu"),
                "ders_adi": course.get("ders_adi"),
                "toplam_ogrenci": None,
                "gecen_ogrenci": None,
                "basari_ortalamasi": None,
                "kontenjan": None,
                "kayitli_ogrenci": None,
                "aciklama": None,
            }
            for course in template_courses
        ]
    else:
        rows = [
            {
                "fakulte_adi": faculty_text,
                "bolum_adi": department_text or FACULTY_SCOPE_LABEL,
                "yil": year_value,
                "donem": term_value,
                "ders_kodu": "SEC101",
                "ders_adi": "Ornek Secmeli Ders",
                "toplam_ogrenci": None,
                "gecen_ogrenci": None,
                "basari_ortalamasi": None,
                "kontenjan": None,
                "kayitli_ogrenci": None,
                "aciklama": None,
            }
        ]

    meta_df = pd.DataFrame(
        [
            {
                "fakulte_adi": faculty_text,
                "bolum_adi": department_text or FACULTY_SCOPE_LABEL,
                "yil": year_value,
                "donem": term_value,
                "aciklama": None,
            }
        ]
    )
    data_df = pd.DataFrame(rows)

    with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
        meta_df.to_excel(writer, sheet_name="Meta", index=False)
        data_df.to_excel(writer, sheet_name=CRITERIA_TEMPLATE_SHEET_NAME, index=False)
        worksheet = writer.sheets[CRITERIA_TEMPLATE_SHEET_NAME]
        for col_idx, header in enumerate(list(data_df.columns), start=1):
            worksheet.cell(row=1, column=col_idx).font = Font(bold=True)
            width = max(len(str(header)), 14)
            letter = _excel_column_letter(col_idx)
            worksheet.column_dimensions[letter].width = width + 2
    return target_path


def _load_scope_candidates(
    cur: sqlite3.Cursor,
    faculty_id: int,
    year: int,
    term: str,
    department_id: int | None = None,
) -> tuple[list[CourseCandidate], list[dict[str, Any]]]:
    scope_courses = _get_scope_courses(
        cur=cur,
        faculty_id=int(faculty_id),
        year=int(year),
        term=normalize_term_label(term),
        department_id=int(department_id) if department_id is not None else None,
    )
    scope_ids = {int(item["ders_id"]) for item in scope_courses if item.get("ders_id") is not None}
    if not scope_ids:
        return [], []

    faculty_candidates = load_faculty_course_candidates(cur=cur, faculty_id=int(faculty_id), year=int(year))
    filtered = [candidate for candidate in faculty_candidates if int(candidate.ders_id) in scope_ids]
    if filtered:
        return filtered, scope_courses

    return (
        [
            CourseCandidate(
                ders_id=int(item["ders_id"]),
                ders_kodu=str(item.get("ders_kodu") or ""),
                ders_adi=str(item.get("ders_adi") or ""),
                in_year_scope=True,
            )
            for item in scope_courses
        ],
        scope_courses,
    )


def match_criteria_rows(
    conn: sqlite3.Connection,
    rows: list[CriteriaRow],
    faculty_id: int,
    year: int,
    term: str,
    department_id: int | None = None,
) -> dict[str, Any]:
    cur = conn.cursor()
    candidates, scope_courses = _load_scope_candidates(
        cur=cur,
        faculty_id=int(faculty_id),
        year=int(year),
        term=normalize_term_label(term),
        department_id=int(department_id) if department_id is not None else None,
    )
    if not candidates:
        return {
            "ok": False,
            "matched_rows": [],
            "unmatched_rows": [],
            "matched_count": 0,
            "unmatched_count": 0,
            "errors": ["Secili kapsam icin eslenecek ders adayi bulunamadi."],
            "scope_courses": scope_courses,
        }

    matched_rows: list[CriteriaImportRowResult] = []
    unmatched_rows: list[CriteriaImportRowResult] = []
    seen_course_ids: dict[int, int] = {}
    errors: list[str] = []

    for row in rows:
        result = match_course_row(candidates=candidates, ders_kodu=row.ders_kodu, ders_adi=row.ders_adi)
        row_result = CriteriaImportRowResult(
            row_no=row.row_no,
            ders_kodu=row.ders_kodu,
            ders_adi=row.ders_adi,
            toplam_ogrenci=int(row.toplam_ogrenci or 0),
            gecen_ogrenci=int(row.gecen_ogrenci or 0),
            basari_ortalamasi=float(row.basari_ortalamasi or 0.0),
            kontenjan=int(row.kontenjan or 0),
            kayitli_ogrenci=int(row.kayitli_ogrenci or 0),
            aciklama=row.aciklama,
            raw_fakulte=row.fakulte_adi,
            raw_bolum=row.bolum_adi,
            raw_yil=row.yil,
            raw_donem=row.donem,
        )
        if not result.matched or result.ders_id is None:
            row_result.row_status = "unmatched"
            row_result.error_message = result.error or "Sistemde eslesen ders bulunamadi."
            unmatched_rows.append(row_result)
            continue

        if int(result.ders_id) in seen_course_ids:
            first_row = seen_course_ids[int(result.ders_id)]
            row_result.row_status = "duplicate"
            row_result.error_message = (
                f"Ayni ders belge icinde birden fazla kez eslesti: satir {first_row} ve {row.row_no}."
            )
            unmatched_rows.append(row_result)
            continue

        seen_course_ids[int(result.ders_id)] = row.row_no
        row_result.matched_ders_id = int(result.ders_id)
        row_result.match_method = result.match_method
        matched_rows.append(row_result)

    if unmatched_rows:
        errors.extend(
            [
                f"Satir {row.row_no}: {row.error_message or 'Sistemde eslesen ders bulunamadi.'}"
                for row in unmatched_rows
            ]
        )

    return {
        "ok": len(errors) == 0,
        "matched_rows": matched_rows,
        "unmatched_rows": unmatched_rows,
        "matched_count": len(matched_rows),
        "unmatched_count": len(unmatched_rows),
        "errors": errors,
        "scope_courses": scope_courses,
    }


def _iter_chunks(values: list[int], chunk_size: int = 800) -> list[list[int]]:
    return [values[idx : idx + chunk_size] for idx in range(0, len(values), chunk_size)]


def _exact_scope_import_ids(
    cur: sqlite3.Cursor,
    faculty_id: int,
    year: int,
    term: str,
    department_id: int | None = None,
    only_applied: bool = True,
) -> list[int]:
    params: list[Any] = [int(faculty_id), int(year), term_key(term)]
    where_status = "AND status = 'applied'" if only_applied else ""
    if department_id is None:
        cur.execute(
            f"""
            SELECT import_id
            FROM criteria_import
            WHERE fakulte_id = ?
              AND bolum_id IS NULL
              AND yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
              {where_status}
            ORDER BY version DESC, import_id DESC
            """,
            tuple(params),
        )
    else:
        params.append(int(department_id))
        cur.execute(
            f"""
            SELECT import_id
            FROM criteria_import
            WHERE fakulte_id = ?
              AND yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
              AND bolum_id = ?
              {where_status}
            ORDER BY version DESC, import_id DESC
            """,
            tuple(params),
        )
    return [int(row[0]) for row in cur.fetchall() if row and row[0] is not None]


def _delete_metrics_for_courses(
    cur: sqlite3.Cursor,
    course_ids: set[int],
    year: int,
    term: str,
) -> dict[str, int]:
    stats = {"performance_rows_deleted": 0, "popularity_rows_deleted": 0}
    if not course_ids:
        return stats
    values = sorted(int(course_id) for course_id in course_ids)
    for chunk in _iter_chunks(values):
        placeholders = ",".join("?" for _ in chunk)
        cur.execute(
            f"""
            DELETE FROM performans
            WHERE akademik_yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
              AND ders_id IN ({placeholders})
            """,
            (int(year), term_key(term), *chunk),
        )
        stats["performance_rows_deleted"] += int(cur.rowcount or 0)
        cur.execute(
            f"""
            DELETE FROM populerlik
            WHERE akademik_yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
              AND ders_id IN ({placeholders})
            """,
            (int(year), term_key(term), *chunk),
        )
        stats["popularity_rows_deleted"] += int(cur.rowcount or 0)
    return stats


def replace_existing_criteria_scope(
    conn: sqlite3.Connection,
    faculty_id: int,
    year: int,
    term: str,
    department_id: int | None = None,
) -> dict[str, int]:
    ensure_criteria_import_schema(conn, commit=False)
    cur = conn.cursor()
    import_ids = _exact_scope_import_ids(
        cur=cur,
        faculty_id=int(faculty_id),
        year=int(year),
        term=normalize_term_label(term),
        department_id=int(department_id) if department_id is not None else None,
        only_applied=True,
    )
    if not import_ids:
        return {
            "previous_imports_superseded": 0,
            "criteria_rows_reset": 0,
            "performance_rows_deleted": 0,
            "popularity_rows_deleted": 0,
        }

    course_ids: set[int] = set()
    for import_id in import_ids:
        cur.execute("SELECT matched_ders_id FROM criteria_import_rows WHERE import_id = ?", (int(import_id),))
        course_ids.update(int(row[0]) for row in cur.fetchall() if row and row[0] is not None)
        cur.execute(
            """
            SELECT ders_id
            FROM ders_kriterleri
            WHERE criteria_import_id = ?
              AND yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
            """,
            (int(import_id), int(year), term_key(term)),
        )
        course_ids.update(int(row[0]) for row in cur.fetchall() if row and row[0] is not None)

    protected_course_ids = (
        _active_department_override_course_ids(cur=cur, course_ids=course_ids, year=int(year), term=term)
        if department_id is None
        else set()
    )

    rows_reset = 0
    for import_id in import_ids:
        cur.execute(
            """
            UPDATE ders_kriterleri
            SET toplam_ogrenci = 0,
                gecen_ogrenci = 0,
                basari_ortalamasi = 0.0,
                kontenjan = 0,
                kayitli_ogrenci = 0,
                criteria_import_id = NULL,
                criteria_veri_kaynagi = 'manual',
                criteria_manual_override = 0,
                criteria_updated_at = ?
            WHERE criteria_import_id = ?
              AND yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
            """,
            (_now_utc(), int(import_id), int(year), term_key(term)),
        )
        rows_reset += int(cur.rowcount or 0)
    metrics_stats = _delete_metrics_for_courses(
        cur=cur,
        course_ids=(course_ids - protected_course_ids),
        year=int(year),
        term=term,
    )

    if import_ids:
        placeholders = ",".join("?" for _ in import_ids)
        cur.execute(
            f"""
            UPDATE criteria_import
            SET status = 'superseded'
            WHERE import_id IN ({placeholders})
            """,
            tuple(int(import_id) for import_id in import_ids),
        )

    return {
        "previous_imports_superseded": len(import_ids),
        "criteria_rows_reset": rows_reset,
        **metrics_stats,
    }


def _next_scope_version(
    cur: sqlite3.Cursor,
    faculty_id: int,
    year: int,
    term: str,
    department_id: int | None = None,
) -> int:
    params: list[Any] = [int(faculty_id), int(year), term_key(term)]
    if department_id is None:
        cur.execute(
            """
            SELECT COALESCE(MAX(version), 0)
            FROM criteria_import
            WHERE fakulte_id = ?
              AND bolum_id IS NULL
              AND yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
            """,
            tuple(params),
        )
    else:
        params.append(int(department_id))
        cur.execute(
            """
            SELECT COALESCE(MAX(version), 0)
            FROM criteria_import
            WHERE fakulte_id = ?
              AND yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
              AND bolum_id = ?
            """,
            tuple(params),
        )
    row = cur.fetchone()
    return int((row[0] if row else 0) or 0) + 1


def _find_criteria_row_id(
    cur: sqlite3.Cursor,
    ders_id: int,
    year: int,
    term: str,
) -> int | None:
    cur.execute(
        """
        SELECT id
        FROM ders_kriterleri
        WHERE ders_id = ?
          AND yil = ?
          AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
        ORDER BY id DESC
        LIMIT 1
        """,
        (int(ders_id), int(year), term_key(term)),
    )
    row = cur.fetchone()
    return int(row[0]) if row and row[0] is not None else None


def _row_has_active_department_override(cur: sqlite3.Cursor, criteria_row_id: int) -> bool:
    cur.execute(
        """
        SELECT ci.import_id
        FROM ders_kriterleri dk
        JOIN criteria_import ci ON ci.import_id = dk.criteria_import_id
        WHERE dk.id = ?
          AND ci.status = 'applied'
          AND ci.bolum_id IS NOT NULL
        LIMIT 1
        """,
        (int(criteria_row_id),),
    )
    return cur.fetchone() is not None


def _active_department_override_course_ids(
    cur: sqlite3.Cursor,
    course_ids: set[int],
    year: int,
    term: str,
) -> set[int]:
    if not course_ids:
        return set()
    protected: set[int] = set()
    for chunk in _iter_chunks(sorted(int(course_id) for course_id in course_ids)):
        placeholders = ",".join("?" for _ in chunk)
        cur.execute(
            f"""
            SELECT DISTINCT dk.ders_id
            FROM ders_kriterleri dk
            JOIN criteria_import ci ON ci.import_id = dk.criteria_import_id
            WHERE dk.yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(dk.donem, '')), 1, 1)) = ?
              AND dk.ders_id IN ({placeholders})
              AND ci.status = 'applied'
              AND ci.bolum_id IS NOT NULL
            """,
            (int(year), term_key(term), *chunk),
        )
        protected.update(int(row[0]) for row in cur.fetchall() if row and row[0] is not None)
    return protected


def apply_criteria_import(
    conn: sqlite3.Connection,
    faculty_id: int,
    year: int,
    term: str,
    rows: list[CriteriaImportRowResult],
    source_filename: str | None = None,
    department_id: int | None = None,
    template_version: str = CRITERIA_TEMPLATE_VERSION,
    notes: str | None = None,
    import_batch_id: int | None = None,
    apply_to_live: bool = True,
) -> dict[str, Any]:
    """Kriter satirlarini stage eder ve (apply_to_live=True ise) canli tablolara uygular.

    apply_to_live=False (ERTELEME): yalniz onizleme/kalite icin criteria_import_rows'a
    stage eder; ders_kriterleri/performans/populerlik'e DOKUNMAZ ve kapsam SIFIRLANMAZ.
    Boylece kriter importu onaylanana kadar algoritmalari etkilemez (§5). Onayda bu
    fonksiyon yeniden (apply_to_live=True ile) cagrilarak canli uygulama yapilir.
    """
    ensure_criteria_import_schema(conn, commit=False)
    cur = conn.cursor()
    normalized_term = normalize_term_label(term)
    now = _now_utc()

    if apply_to_live:
        replace_stats = replace_existing_criteria_scope(
            conn=conn,
            faculty_id=int(faculty_id),
            year=int(year),
            term=normalized_term,
            department_id=int(department_id) if department_id is not None else None,
        )
    else:
        replace_stats = {}

    version = _next_scope_version(
        cur=cur,
        faculty_id=int(faculty_id),
        year=int(year),
        term=normalized_term,
        department_id=int(department_id) if department_id is not None else None,
    )
    criteria_status = "applied" if apply_to_live else "staged"
    cur.execute(
        """
        INSERT INTO criteria_import
            (fakulte_id, bolum_id, yil, donem, source_filename, template_version, notes, imported_at, status, version)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            int(faculty_id),
            int(department_id) if department_id is not None else None,
            int(year),
            normalized_term,
            source_filename,
            template_version,
            notes,
            now,
            criteria_status,
            int(version),
        ),
    )
    import_id = int(cur.lastrowid or 0)
    if import_batch_id is not None:
        cur.execute(
            """
            UPDATE criteria_import
            SET import_batch_id = ?
            WHERE import_id = ?
            """,
            (int(import_batch_id), int(import_id)),
        )

    created_rows = 0
    updated_rows = 0
    skipped_department_overrides = 0
    applied_course_ids: set[int] = set()

    for row in rows:
        cur.execute(
            """
            INSERT INTO criteria_import_rows
                (import_id, row_no, ders_kodu, ders_adi, toplam_ogrenci, gecen_ogrenci,
                 basari_ortalamasi, kontenjan, kayitli_ogrenci, matched_ders_id, match_method,
                 row_status, error_message, raw_fakulte, raw_bolum, raw_yil, raw_donem)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(import_id),
                int(row.row_no),
                row.ders_kodu,
                row.ders_adi,
                int(row.toplam_ogrenci),
                int(row.gecen_ogrenci),
                float(row.basari_ortalamasi),
                int(row.kontenjan),
                int(row.kayitli_ogrenci),
                int(row.matched_ders_id) if row.matched_ders_id is not None else None,
                row.match_method,
                row.row_status,
                row.error_message,
                row.raw_fakulte,
                row.raw_bolum,
                row.raw_yil,
                row.raw_donem,
            ),
        )
        import_row_id = int(cur.lastrowid or 0)
        if import_batch_id is not None:
            normalized_row = asdict(row)
            cur.execute(
                """
                UPDATE criteria_import_rows
                SET import_batch_id = ?, normalized_row_json = ?, row_hash = ?
                WHERE row_id = ?
                """,
                (
                    int(import_batch_id),
                    json.dumps(normalized_row, ensure_ascii=False, sort_keys=True, default=str),
                    calculate_row_hash(normalized_row),
                    int(import_row_id),
                ),
            )
            if row.error_message:
                record_import_issue(
                    conn,
                    import_batch_id=int(import_batch_id),
                    source_row_id=int(import_row_id),
                    row_number=int(row.row_no),
                    message=str(row.error_message),
                )
                cur.execute(
                    "UPDATE criteria_import_rows SET issue_count = COALESCE(issue_count, 0) + 1 WHERE row_id = ?",
                    (int(import_row_id),),
                )

        # ERTELEME: apply_to_live=False ise satir stage edildi; canli tablolara
        # (ders_kriterleri/performans/populerlik) yazma yapilmaz. Onayda uygulanir.
        if not apply_to_live:
            continue

        if row.matched_ders_id is None:
            continue

        existing_id = _find_criteria_row_id(
            cur=cur,
            ders_id=int(row.matched_ders_id),
            year=int(year),
            term=normalized_term,
        )
        if department_id is None and existing_id is not None and _row_has_active_department_override(cur, existing_id):
            skipped_department_overrides += 1
            cur.execute(
                """
                UPDATE criteria_import_rows
                SET row_status = 'skipped_override',
                    error_message = 'Daha ozel bolum kriter dosyasi aktif oldugu icin fakulte geneli veri uygulanmadi.'
                WHERE row_id = ?
                """,
                (int(import_row_id),),
            )
            if import_batch_id is not None:
                record_import_issue(
                    conn,
                    import_batch_id=int(import_batch_id),
                    source_row_id=int(import_row_id),
                    row_number=int(row.row_no),
                    severity="info",
                    issue_type="invalid_scope",
                    message="Daha ozel bolum kriter dosyasi aktif oldugu icin fakulte geneli veri uygulanmadi.",
                    suggestion="Bu ders icin bolum ozel import aktif oldugundan fakulte geneli dosyada islem yapmaya gerek yoktur.",
                )
                cur.execute(
                    "UPDATE criteria_import_rows SET issue_count = COALESCE(issue_count, 0) + 1 WHERE row_id = ?",
                    (int(import_row_id),),
                )
            continue

        if existing_id is not None:
            cur.execute(
                """
                UPDATE ders_kriterleri
                SET toplam_ogrenci = ?,
                    gecen_ogrenci = ?,
                    basari_ortalamasi = ?,
                    kontenjan = ?,
                    kayitli_ogrenci = ?,
                    criteria_import_id = ?,
                    criteria_veri_kaynagi = 'criteria_import',
                    criteria_manual_override = 0,
                    criteria_updated_at = ?
                WHERE id = ?
                """,
                (
                    int(row.toplam_ogrenci),
                    int(row.gecen_ogrenci),
                    float(row.basari_ortalamasi),
                    int(row.kontenjan),
                    int(row.kayitli_ogrenci),
                    int(import_id),
                    now,
                    int(existing_id),
                ),
            )
            updated_rows += int(cur.rowcount or 0)
        else:
            cur.execute(
                """
                INSERT INTO ders_kriterleri
                    (ders_id, yil, donem, toplam_ogrenci, gecen_ogrenci, basari_ortalamasi,
                     kontenjan, kayitli_ogrenci, anket_katilimci, anket_dersi_secen,
                     anket_veri_kaynagi, anket_manual_locked, anket_import_id, anket_imported_at,
                     criteria_import_id, criteria_veri_kaynagi, criteria_manual_override, criteria_updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, 0, 0, 'manual', 0, NULL, NULL, ?, 'criteria_import', 0, ?)
                """,
                (
                    int(row.matched_ders_id),
                    int(year),
                    normalized_term,
                    int(row.toplam_ogrenci),
                    int(row.gecen_ogrenci),
                    float(row.basari_ortalamasi),
                    int(row.kontenjan),
                    int(row.kayitli_ogrenci),
                    int(import_id),
                    now,
                ),
            )
            created_rows += 1

        basari_orani = (float(row.gecen_ogrenci) / float(row.toplam_ogrenci)) if row.toplam_ogrenci > 0 else 0.0
        doluluk_orani = (
            min(float(row.kayitli_ogrenci) / float(row.kontenjan), 1.0) if row.kontenjan > 0 else 0.0
        )

        cur.execute(
            """
            DELETE FROM performans
            WHERE ders_id = ?
              AND akademik_yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
            """,
            (int(row.matched_ders_id), int(year), term_key(normalized_term)),
        )
        cur.execute(
            """
            INSERT INTO performans (ders_id, akademik_yil, donem, ortalama_not, basari_orani)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                int(row.matched_ders_id),
                int(year),
                normalized_term,
                float(row.basari_ortalamasi),
                float(basari_orani),
            ),
        )
        cur.execute(
            """
            DELETE FROM populerlik
            WHERE ders_id = ?
              AND akademik_yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
            """,
            (int(row.matched_ders_id), int(year), term_key(normalized_term)),
        )
        cur.execute(
            """
            INSERT INTO populerlik (ders_id, akademik_yil, donem, talep_sayisi, kontenjan, doluluk_orani)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                int(row.matched_ders_id),
                int(year),
                normalized_term,
                int(row.kayitli_ogrenci),
                int(row.kontenjan),
                float(doluluk_orani),
            ),
        )
        applied_course_ids.add(int(row.matched_ders_id))
        if import_batch_id is not None:
            for field_name, value in (
                ("toplam_ogrenci", row.toplam_ogrenci),
                ("gecen_ogrenci", row.gecen_ogrenci),
                ("basari_ortalamasi", row.basari_ortalamasi),
                ("kontenjan", row.kontenjan),
                ("kayitli_ogrenci", row.kayitli_ogrenci),
            ):
                record_value_source(
                    conn=conn,
                    course_id=int(row.matched_ders_id),
                    year=int(year),
                    field_name=field_name,
                    value=value,
                    source_type="criteria_import",
                    faculty_id=int(faculty_id),
                    department_id=int(department_id) if department_id is not None else None,
                    source_import_batch_id=int(import_batch_id),
                    source_row_id=int(import_row_id),
                    is_locked=False,
                    deactivate_existing=True,
                )

    if apply_to_live:
        status_result = mark_criteria_status(
            conn=conn,
            yil=int(year),
            fakulte_id=int(faculty_id),
            bolum_id=int(department_id) if department_id is not None else None,
        )
    else:
        status_result = {}

    return {
        "ok": True,
        "import_id": import_id,
        "version": version,
        "created_rows": created_rows,
        "updated_rows": updated_rows,
        "skipped_department_overrides": skipped_department_overrides,
        "applied_course_count": len(applied_course_ids),
        "replace": replace_stats,
        "status": status_result,
        "staged": not apply_to_live,
    }


def import_criteria_excel(
    db_path: str,
    excel_path: str,
    faculty_id: int,
    year: int,
    term: str,
    department_id: int | None = None,
    source_filename: str | None = None,
    auto_activate: bool = True,
    uploaded_by: str | None = None,
    apply_now: bool = True,
) -> dict[str, Any]:
    resolved_db_path = resolve_sqlite_db_path(db_path)
    if not resolved_db_path.exists():
        return {"ok": False, "message": "Veritabani bulunamadi.", "errors": ["Veritabani bulunamadi."]}
    if not os.path.exists(excel_path):
        return {"ok": False, "message": "Kriter dosyasi bulunamadi.", "errors": ["Kriter dosyasi bulunamadi."]}

    try:
        parsed = parse_criteria_excel(excel_path)
    except Exception as exc:
        try:
            mark_batch_failed_by_path(
                str(resolved_db_path),
                excel_path,
                "criteria",
                str(exc),
                original_filename=source_filename or os.path.basename(excel_path),
                faculty_id=int(faculty_id),
                department_id=int(department_id) if department_id is not None else None,
                year=int(year),
                semester=normalize_term_label(term),
            )
        except Exception:
            pass
        return {"ok": False, "message": f"Kriter dosyasi okunamadi: {exc}", "errors": [str(exc)]}

    conn = connect_sqlite(str(resolved_db_path), row_factory=True)
    try:
        ensure_reporting_schema(conn)
        ensure_criteria_import_schema(conn)
        cur = conn.cursor()
        try:
            metadata = extract_excel_metadata(excel_path)
        except Exception:
            metadata = {
                "sheet_names": [str(parsed.get("sheet_name") or "Kriter")],
                "columns": [],
                "row_count": len(parsed.get("rows") or []),
                "column_count": 0,
            }
        batch = create_import_batch(
            conn,
            import_type="criteria",
            original_filename=source_filename or os.path.basename(excel_path),
            file_path=excel_path,
            sheet_names=list(metadata.get("sheet_names") or []),
            columns=list(metadata.get("columns") or []),
            row_count=int(metadata.get("row_count") or len(parsed.get("rows") or [])),
            column_count=int(metadata.get("column_count") or 0),
            faculty_id=int(faculty_id),
            department_id=int(department_id) if department_id is not None else None,
            year=int(year),
            semester=normalize_term_label(term),
            uploaded_by=uploaded_by,
            status="uploaded",
        )
        import_batch_id = int(batch["import_batch_id"])
        conn.commit()
        faculty_name = _resolve_faculty_name(cur, int(faculty_id))
        if not faculty_name:
            record_import_issue(
                conn,
                import_batch_id=import_batch_id,
                row_number=0,
                severity="critical",
                issue_type="invalid_scope",
                message="Secili fakulte bulunamadi.",
                suggestion="Import kapsaminda gecerli bir fakulte secin.",
            )
            update_import_status(conn, import_batch_id, "failed", error_message="Secili fakulte bulunamadi.")
            conn.commit()
            return {"ok": False, "message": "Secili fakulte bulunamadi.", "errors": ["Secili fakulte bulunamadi."]}
        department_name = _resolve_department_name(cur, int(department_id)) if department_id is not None else None

        validation = validate_criteria_rows(
            rows=list(parsed.get("rows") or []),
            faculty_name=faculty_name,
            department_name=department_name,
            year=int(year),
            term=normalize_term_label(term),
        )
        warnings = list(parsed.get("warnings") or []) + list(validation.get("warnings") or [])
        if not validation.get("ok"):
            for error in list(validation.get("errors") or []):
                record_import_issue(
                    conn,
                    import_batch_id=import_batch_id,
                    row_number=0,
                    message=str(error),
                )
            quality = evaluate_import_quality(conn, import_batch_id)
            update_import_status(
                conn,
                import_batch_id,
                "failed",
                error_message="Kriter belgesi dogrulamasi basarisiz.",
                validation_summary={"ok": False, "errors": list(validation.get("errors") or []), "warnings": warnings},
            )
            conn.commit()
            return {
                "ok": False,
                "message": "Kriter belgesi dogrulamasi basarisiz.",
                "errors": list(validation.get("errors") or []),
                "warnings": warnings,
                "matched_count": 0,
                "unmatched_count": 0,
                "import_batch_id": import_batch_id,
                "quality_score": quality.quality_score,
                "quality_level": quality.quality_level,
            }

        matched = match_criteria_rows(
            conn=conn,
            rows=list(parsed.get("rows") or []),
            faculty_id=int(faculty_id),
            year=int(year),
            term=normalize_term_label(term),
            department_id=int(department_id) if department_id is not None else None,
        )
        if not matched.get("ok"):
            for unmatched in matched.get("unmatched_rows") or []:
                row_dict = unmatched.as_dict() if hasattr(unmatched, "as_dict") else {}
                record_import_issue(
                    conn,
                    import_batch_id=import_batch_id,
                    row_number=int(row_dict.get("row_no") or 0),
                    issue_type="course_not_matched",
                    severity="error",
                    message=row_dict.get("error_message") or "Ders secili kapsamda eslesmedi.",
                    suggestion="Ders kodu veya adini sistemdeki ders kaydi ile uyumlu hale getirin.",
                )
            quality = evaluate_import_quality(conn, import_batch_id)
            update_import_status(
                conn,
                import_batch_id,
                "pending_review",
                error_message="Belgedeki bazi dersler secili kapsamda eslesmedi.",
                validation_summary={
                    "ok": False,
                    "matched_count": int(matched.get("matched_count") or 0),
                    "unmatched_count": int(matched.get("unmatched_count") or 0),
                },
            )
            conn.commit()
            return {
                "ok": False,
                "message": "Belgedeki bazi dersler secili kapsamda eslesmedi. Veri uygulanmadi.",
                "errors": list(matched.get("errors") or []),
                "warnings": warnings,
                "matched_count": int(matched.get("matched_count") or 0),
                "unmatched_count": int(matched.get("unmatched_count") or 0),
                "matched_rows": [row.as_dict() for row in matched.get("matched_rows") or []],
                "unmatched_rows": [row.as_dict() for row in matched.get("unmatched_rows") or []],
                "import_batch_id": import_batch_id,
                "quality_score": quality.quality_score,
                "quality_level": quality.quality_level,
            }

        conn.execute("BEGIN")
        applied = apply_criteria_import(
            conn=conn,
            faculty_id=int(faculty_id),
            year=int(year),
            term=normalize_term_label(term),
            rows=list(matched.get("matched_rows") or []),
            source_filename=source_filename or os.path.basename(excel_path),
            department_id=int(department_id) if department_id is not None else None,
            template_version=str(parsed.get("template_version") or CRITERIA_TEMPLATE_VERSION),
            notes=(parsed.get("meta") or {}).get("aciklama"),
            import_batch_id=import_batch_id,
            apply_to_live=apply_now,
        )
        link_source_import(
            conn,
            import_batch_id=import_batch_id,
            source_table="criteria_import",
            source_import_id=int(applied["import_id"]),
            file_hash_sha256=batch.get("file_hash_sha256"),
            file_size=batch.get("file_size"),
        )
        quality = evaluate_import_quality(conn, import_batch_id)
        # ERTELEME (§5): apply_now=False ise kriterler henuz uygulanmadi; import
        # onay bekler. Aksi halde mevcut davranis (dusuk kalite -> pending).
        if not apply_now:
            status = "pending_review"
        else:
            status = "pending_review" if quality.quality_level == "low" else "validated"
        update_import_status(
            conn,
            import_batch_id,
            status,
            validation_summary={
                "ok": True,
                "matched_count": int(matched.get("matched_count") or 0),
                "unmatched_count": int(matched.get("unmatched_count") or 0),
                "warnings": warnings,
                "staged": not apply_now,
            },
        )
        if apply_now and auto_activate and quality.quality_level != "low":
            from app.services.import_audit_service import activate_import

            activate_import(conn, import_batch_id, user=uploaded_by)
        try:
            recalculate_import_diff(conn, import_batch_id)
        except Exception:
            pass
        try:
            recalculate_import_impact(conn, import_batch_id)
        except Exception:
            pass
        conn.commit()

        if apply_now:
            message = (
                "Kriter belgesi basariyla yuklendi. "
                f"Versiyon: v{applied['version']} | Uygulanan ders: {applied['applied_course_count']}."
            )
        else:
            message = (
                f"Kriter belgesi yuklendi ve ONAY BEKLIYOR (eslesen ders: {int(matched.get('matched_count') or 0)}). "
                "Kriterler henuz uygulanmadi; 'Rollback & Onay' sekmesinden 'Aktif Yap' ile uygulayin."
            )
        return {
            "ok": True,
            "message": message,
            "staged": not apply_now,
            "faculty_id": int(faculty_id),
            "department_id": int(department_id) if department_id is not None else None,
            "year": int(year),
            "term": normalize_term_label(term),
            "matched_count": int(matched.get("matched_count") or 0),
            "unmatched_count": int(matched.get("unmatched_count") or 0),
            "updated_course_count": int(applied.get("updated_rows") or 0),
            "created_course_count": int(applied.get("created_rows") or 0),
            "applied_course_count": int(applied.get("applied_course_count") or 0),
            "skipped_department_overrides": int(applied.get("skipped_department_overrides") or 0),
            "replace": dict(applied.get("replace") or {}),
            "status": dict(applied.get("status") or {}),
            "warnings": warnings,
            "matched_rows": [row.as_dict() for row in matched.get("matched_rows") or []],
            "unmatched_rows": [],
            "import_id": int(applied["import_id"]),
            "import_batch_id": import_batch_id,
            "import_status": "active" if apply_now and auto_activate and quality.quality_level != "low" else status,
            "quality_score": quality.quality_score,
            "quality_level": quality.quality_level,
            "duplicate": bool(batch.get("duplicate")),
            "duplicate_of_import_batch_id": batch.get("duplicate_of_import_batch_id"),
            "version": int(applied["version"]),
        }
    except Exception as exc:
        conn.rollback()
        try:
            failed_batch_id = locals().get("import_batch_id")
            if failed_batch_id is not None:
                update_import_status(conn, int(failed_batch_id), "failed", error_message=str(exc))
                conn.commit()
        except Exception:
            pass
        return {
            "ok": False,
            "message": f"Kriter yukleme hatasi: {exc}",
            "errors": [str(exc)],
            "warnings": list(parsed.get("warnings") or []),
        }
    finally:
        conn.close()


def apply_pending_criteria_import(
    conn: sqlite3.Connection,
    import_batch_id: int,
    user: str | None = None,
) -> dict[str, Any]:
    """§5: Onay bekleyen (staged) bir kriter importunu canli tablolara uygular.

    import sirasinda ``apply_to_live=False`` ile yalniz stage edilen satirlari
    (``criteria_import_rows.normalized_row_json``) yeniden kurar ve
    ``apply_criteria_import``'u ``apply_to_live=True`` ile cagirarak
    ders_kriterleri/performans/populerlik'e uygular; ardindan importu aktive eder.
    Caller commit eder. Staged satir yoksa ``ok=False`` doner (UI bunu activate'e cevirir).
    """
    from app.services.import_audit_service import activate_import, get_import_batch

    ensure_criteria_import_schema(conn, commit=False)
    batch = get_import_batch(conn, int(import_batch_id))
    if not batch:
        return {"ok": False, "message": "Import kaydi bulunamadi."}
    if str(batch.get("import_type")) != "criteria":
        return {"ok": False, "message": "Bu islem yalniz kriter importlari icindir."}
    if str(batch.get("status") or "").lower() in ("active", "applied", "rejected", "rolled_back"):
        return {"ok": False, "message": "Import zaten uygulanmis/aktif veya kapali."}

    cur = conn.cursor()
    cur.execute(
        "SELECT normalized_row_json FROM criteria_import_rows WHERE import_batch_id = ? ORDER BY row_no",
        (int(import_batch_id),),
    )
    staged_payloads = [r[0] for r in cur.fetchall() if r and r[0]]
    if not staged_payloads:
        return {"ok": False, "message": "Uygulanacak stage edilmis kriter satiri yok."}

    rows: list[CriteriaImportRowResult] = []
    for payload in staged_payloads:
        try:
            rows.append(CriteriaImportRowResult(**json.loads(payload)))
        except Exception:
            continue
    if not rows:
        return {"ok": False, "message": "Stage edilmis satirlar cozumlenemedi."}

    faculty_id = int(batch.get("faculty_id"))
    department_id = batch.get("department_id")
    year = int(batch.get("year"))
    term = normalize_term_label(batch.get("semester") or "Guz")

    # Eski staged kayitlari temizle; apply yeniden temiz stage edip uygulasin (kopya olmasin).
    cur.execute("DELETE FROM criteria_import_rows WHERE import_batch_id = ?", (int(import_batch_id),))
    cur.execute("DELETE FROM criteria_import WHERE import_batch_id = ?", (int(import_batch_id),))

    applied = apply_criteria_import(
        conn=conn,
        faculty_id=faculty_id,
        year=year,
        term=term,
        rows=rows,
        source_filename=batch.get("original_filename"),
        department_id=int(department_id) if department_id is not None else None,
        import_batch_id=int(import_batch_id),
        apply_to_live=True,
    )
    link_source_import(
        conn,
        import_batch_id=int(import_batch_id),
        source_table="criteria_import",
        source_import_id=int(applied["import_id"]),
        file_hash_sha256=batch.get("file_hash_sha256"),
        file_size=batch.get("file_size"),
    )
    activate_import(conn, int(import_batch_id), user=user)
    try:
        recalculate_import_impact(conn, int(import_batch_id))
    except Exception:
        pass
    return {
        "ok": True,
        "message": f"Kriterler uygulandi ve aktive edildi. Uygulanan ders: {int(applied.get('applied_course_count') or 0)}.",
        "applied_course_count": int(applied.get("applied_course_count") or 0),
        "import_batch_id": int(import_batch_id),
    }


def _criteria_import_row_to_summary(row: sqlite3.Row | tuple[Any, ...] | None) -> dict[str, Any] | None:
    if not row:
        return None
    return {
        "import_id": int(row[0]),
        "fakulte_id": int(row[1]),
        "bolum_id": int(row[2]) if row[2] is not None else None,
        "yil": int(row[3]),
        "donem": normalize_term_label(row[4]),
        "source_filename": str(row[5] or "Bilinmeyen Dosya"),
        "template_version": str(row[6] or ""),
        "notes": row[7],
        "imported_at": row[8],
        "status": str(row[9] or ""),
        "version": int(row[10] or 0),
        "faculty_name": str(row[11] or ""),
        "department_name": str(row[12] or "") or None,
    }


def get_criteria_import_by_id(conn: sqlite3.Connection, import_id: int) -> dict[str, Any] | None:
    ensure_criteria_import_schema(conn, commit=False)
    cur = conn.cursor()
    cur.execute(
        """
        SELECT
            ci.import_id,
            ci.fakulte_id,
            ci.bolum_id,
            ci.yil,
            ci.donem,
            ci.source_filename,
            ci.template_version,
            ci.notes,
            ci.imported_at,
            ci.status,
            ci.version,
            f.ad AS faculty_name,
            b.ad AS department_name
        FROM criteria_import ci
        LEFT JOIN fakulte f ON f.fakulte_id = ci.fakulte_id
        LEFT JOIN bolum b ON b.bolum_id = ci.bolum_id
        WHERE ci.import_id = ?
        LIMIT 1
        """,
        (int(import_id),),
    )
    return _criteria_import_row_to_summary(cur.fetchone())


def get_active_criteria_import(
    conn: sqlite3.Connection,
    faculty_id: int,
    year: int,
    term: str,
    department_id: int | None = None,
) -> dict[str, Any] | None:
    ensure_criteria_import_schema(conn, commit=False)
    cur = conn.cursor()
    params: list[Any] = [int(faculty_id), int(year), term_key(term)]
    if department_id is None:
        cur.execute(
            """
            SELECT
                ci.import_id,
                ci.fakulte_id,
                ci.bolum_id,
                ci.yil,
                ci.donem,
                ci.source_filename,
                ci.template_version,
                ci.notes,
                ci.imported_at,
                ci.status,
                ci.version,
                f.ad AS faculty_name,
                b.ad AS department_name
            FROM criteria_import ci
            LEFT JOIN fakulte f ON f.fakulte_id = ci.fakulte_id
            LEFT JOIN bolum b ON b.bolum_id = ci.bolum_id
            WHERE ci.fakulte_id = ?
              AND ci.bolum_id IS NULL
              AND ci.yil = ?
              AND LOWER(SUBSTR(TRIM(COALESCE(ci.donem, '')), 1, 1)) = ?
              AND ci.status = 'applied'
            ORDER BY ci.version DESC, ci.import_id DESC
            LIMIT 1
            """,
            tuple(params),
        )
        return _criteria_import_row_to_summary(cur.fetchone())

    cur.execute(
        """
        SELECT
            ci.import_id,
            ci.fakulte_id,
            ci.bolum_id,
            ci.yil,
            ci.donem,
            ci.source_filename,
            ci.template_version,
            ci.notes,
            ci.imported_at,
            ci.status,
            ci.version,
            f.ad AS faculty_name,
            b.ad AS department_name
        FROM criteria_import ci
        LEFT JOIN fakulte f ON f.fakulte_id = ci.fakulte_id
        LEFT JOIN bolum b ON b.bolum_id = ci.bolum_id
        WHERE ci.fakulte_id = ?
          AND ci.yil = ?
          AND LOWER(SUBSTR(TRIM(COALESCE(ci.donem, '')), 1, 1)) = ?
          AND ci.status = 'applied'
          AND (ci.bolum_id = ? OR ci.bolum_id IS NULL)
        ORDER BY CASE WHEN ci.bolum_id = ? THEN 0 ELSE 1 END, ci.version DESC, ci.import_id DESC
        LIMIT 1
        """,
        (int(faculty_id), int(year), term_key(term), int(department_id), int(department_id)),
    )
    return _criteria_import_row_to_summary(cur.fetchone())


def format_criteria_import_summary(summary: dict[str, Any] | None) -> str:
    if not summary:
        return "Aktif kriter dosyasi yok."
    scope = summary.get("department_name") or FACULTY_SCOPE_LABEL
    imported_at = summary.get("imported_at") or "-"
    filename = summary.get("source_filename") or "Bilinmeyen Dosya"
    version = int(summary.get("version") or 0)
    return (
        f"{filename} | v{version} | Kapsam: {scope} / {summary.get('yil')} {summary.get('donem')} | "
        f"Yukleme: {imported_at}"
    )


def summarize_report_criteria_scope(
    conn: sqlite3.Connection,
    faculty_id: int,
    year: int,
    term: str,
    department_id: int | None = None,
) -> dict[str, Any]:
    ensure_criteria_import_schema(conn, commit=False)
    if department_id is not None:
        active = get_active_criteria_import(
            conn=conn,
            faculty_id=int(faculty_id),
            year=int(year),
            term=normalize_term_label(term),
            department_id=int(department_id),
        )
        return {
            "mode": "single" if active else "missing",
            "active_import": active,
            "display": format_criteria_import_summary(active),
        }

    cur = conn.cursor()
    cur.execute(
        """
        SELECT COUNT(*)
        FROM criteria_import
        WHERE fakulte_id = ?
          AND bolum_id IS NOT NULL
          AND yil = ?
          AND LOWER(SUBSTR(TRIM(COALESCE(donem, '')), 1, 1)) = ?
          AND status = 'applied'
        """,
        (int(faculty_id), int(year), term_key(term)),
    )
    department_count = int((cur.fetchone() or [0])[0] or 0)
    faculty_import = get_active_criteria_import(
        conn=conn,
        faculty_id=int(faculty_id),
        year=int(year),
        term=normalize_term_label(term),
        department_id=None,
    )
    if department_count > 0:
        if faculty_import:
            display = (
                f"Bu raporda {department_count} bolum bazli aktif kriter dosyasi var. "
                f"Fakulte geneli son dosya: {format_criteria_import_summary(faculty_import)}"
            )
        else:
            display = f"Bu raporda {department_count} bolum bazli aktif kriter dosyasi kullaniliyor."
        return {
            "mode": "mixed",
            "active_import": faculty_import,
            "department_specific_count": department_count,
            "display": display,
        }
    if faculty_import:
        return {
            "mode": "single",
            "active_import": faculty_import,
            "display": format_criteria_import_summary(faculty_import),
        }
    return {
        "mode": "missing",
        "active_import": None,
        "department_specific_count": 0,
        "display": "Aktif kriter dosyasi yok.",
    }
